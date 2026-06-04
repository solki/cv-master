from app.core.settings import get_settings

settings = get_settings()


class ExportRenderer:
    """Renders approved resume versions to Markdown, HTML, PDF, and DOCX."""

    def render_markdown(self, content_json: dict) -> str:
        md = self._build_markdown(content_json)
        return md

    def render_html(self, content_json: dict) -> str:
        md = self._build_markdown(content_json)
        html = self._markdown_to_html(md)
        return self._wrap_html(html, content_json.get("header", {}))

    def render_pdf(self, content_json: dict) -> bytes:
        """Render PDF from HTML content."""
        html = self.render_html(content_json)
        try:
            import weasyprint
            return weasyprint.HTML(string=html).write_pdf()
        except ImportError:
            return html.encode("utf-8")

    def render_docx(self, content_json: dict) -> bytes:
        """Render DOCX from content."""
        md = self._build_markdown(content_json)
        try:
            from docx import Document
            doc = Document()
            for line in md.split("\n"):
                if line.startswith("# "):
                    doc.add_heading(line[2:], level=1)
                elif line.startswith("## "):
                    doc.add_heading(line[3:], level=2)
                elif line.strip().startswith("- "):
                    doc.add_paragraph(line.strip()[2:], style="List Bullet")
                elif line.strip():
                    doc.add_paragraph(line)
            import io
            buf = io.BytesIO()
            doc.save(buf)
            return buf.getvalue()
        except ImportError:
            return md.encode("utf-8")

    def _build_markdown(self, content_json: dict) -> str:
        sections = ["header", "summary", "skills", "experience", "projects", "education", "certifications"]
        lines = []

        header = content_json.get("header", {})
        if header:
            lines.append(f"# {header.get('full_name', '')}")
            lines.append(f"{header.get('email', '')} | {header.get('phone', '')} | {header.get('location', '')}")
            lines.append("")

        summary = content_json.get("summary", "")
        if summary:
            lines.append("## Summary")
            lines.append(str(summary))
            lines.append("")

        skills = content_json.get("skills", [])
        if skills:
            lines.append("## Skills")
            lines.append(", ".join(skills) if isinstance(skills, list) else str(skills))
            lines.append("")

        for section, title in [("experience", "Experience"), ("projects", "Projects"), ("education", "Education"), ("certifications", "Certifications")]:
            entries = content_json.get(section, [])
            if entries:
                lines.append(f"## {title}")
                for entry in entries:
                    if isinstance(entry, dict):
                        role = entry.get("title", entry.get("role", ""))
                        org = entry.get("company", entry.get("organization", entry.get("institution", "")))
                        lines.append(f"### {role} — {org}")
                        desc = entry.get("description", entry.get("summary", ""))
                        if desc:
                            lines.append(str(desc))
                        bullets = entry.get("bullets", [])
                        for b in bullets:
                            lines.append(f"- {b}")
                        lines.append("")
        return "\n".join(lines)

    def _markdown_to_html(self, md: str) -> str:
        try:
            import markdown
            return markdown.markdown(md, extensions=["extra"])
        except ImportError:
            return f"<pre>{md}</pre>"

    def _wrap_html(self, body: str, header: dict) -> str:
        name = header.get("full_name", "CV Master Resume")
        return f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>{name}</title>
    <style>
        body {{ font-family: system-ui, sans-serif; max-width: 800px; margin: 2rem auto; line-height: 1.5; }}
        h1 {{ font-size: 1.5rem; }}
        h2 {{ font-size: 1.2rem; border-bottom: 1px solid #ccc; }}
    </style>
</head>
<body>{body}</body>
</html>"""
